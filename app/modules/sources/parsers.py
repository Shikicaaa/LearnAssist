import re
import zipfile
from dataclasses import dataclass
from enum import StrEnum
from typing import BinaryIO

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from pypdf import PasswordType, PdfReader

from app.modules.jobs import PermanentError

MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MIN_CHARS_PER_PAGE = 25
MAX_EMPTY_PAGE_RATIO = 0.5
TEXT_EXTENSIONS = {".txt": "txt", ".md": "md", ".markdown": "md"}


class SourceFormat(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    TXT = "txt"
    MD = "md"


class ParseError(PermanentError):
    """The file cannot be processed. The message is shown to the user, so keep it readable."""


@dataclass(frozen=True)
class ParsedPage:
    number: int | None  # None when the format has no reliable page numbers
    text: str


@dataclass(frozen=True)
class ParsedDocument:
    pages: list[ParsedPage]
    page_count: int | None


def detect_format(data: BinaryIO, filename: str) -> SourceFormat | None:
    """Detects the format from the file content (magic bytes), not from the extension.

    The extension is only used to tell .txt from .md, since plain text has no signature.
    Leaves the stream rewound to the start.
    """
    head = data.read(4096)
    data.seek(0)
    if not head:
        return None
    if b"%PDF-" in head[:1024]:
        return SourceFormat.PDF
    if head.startswith(b"PK\x03\x04"):
        return SourceFormat.DOCX if _is_docx(data) else None
    if _looks_like_text(head):
        extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        text_format = TEXT_EXTENSIONS.get(extension)
        return SourceFormat(text_format) if text_format else None
    return None


def _is_docx(data: BinaryIO) -> bool:
    try:
        with zipfile.ZipFile(data) as archive:
            names = set(archive.namelist())
        return "[Content_Types].xml" in names and "word/document.xml" in names
    except zipfile.BadZipFile:
        return False
    finally:
        data.seek(0)


def _looks_like_text(head: bytes) -> bool:
    if b"\x00" in head:
        return False
    control = sum(1 for byte in head if byte < 32 and byte not in (9, 10, 13))
    return control / len(head) < 0.01


def parse_document(fmt: SourceFormat, data: BinaryIO, max_pdf_pages: int = 500) -> ParsedDocument:
    if fmt == SourceFormat.PDF:
        return parse_pdf(data, max_pdf_pages)
    if fmt == SourceFormat.DOCX:
        return parse_docx(data)
    return parse_text_file(data)


def _normalize_newlines(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def parse_pdf(data: BinaryIO, max_pages: int = 500) -> ParsedDocument:
    try:
        reader = PdfReader(data)
        # An empty password opens PDFs that only restrict editing/printing (owner password).
        if reader.is_encrypted and reader.decrypt("") == PasswordType.NOT_DECRYPTED:
            raise ParseError("The PDF is password-protected. Remove the password and try again.")
        total = len(reader.pages)
        if total == 0:
            raise ParseError("The PDF has no pages.")
        if total > max_pages:
            raise ParseError(f"The PDF has {total} pages; the limit is {max_pages}.")
        pages: list[ParsedPage] = []
        nearly_empty = 0
        for number, page in enumerate(reader.pages, start=1):
            text = _normalize_newlines(re.sub(r"[ \t]+", " ", page.extract_text() or ""))
            if len(text) < MIN_CHARS_PER_PAGE:
                nearly_empty += 1
            if text:
                pages.append(ParsedPage(number, text))
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("The PDF could not be read; the file may be corrupted.") from exc
    if nearly_empty / total > MAX_EMPTY_PAGE_RATIO:
        raise ParseError(
            "The PDF contains almost no text; it looks like a scan. "
            "Scanned documents (OCR) are not supported yet."
        )
    return ParsedDocument(pages, total)


def parse_docx(data: BinaryIO) -> ParsedDocument:
    try:
        _check_docx_size(data)
        document = Document(data)
        blocks_by_page, has_page_info = _read_docx_body(document)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(
            "The DOCX could not be read; the file may be corrupted or password-protected."
        ) from exc

    if has_page_info:
        pages = [
            ParsedPage(number, "\n\n".join(blocks))
            for number, blocks in enumerate(blocks_by_page, start=1)
            if blocks
        ]
        page_count: int | None = len(blocks_by_page)
    else:
        # Without Word's rendered page markers, page numbers would be guesses, so we give none.
        all_blocks = [block for blocks in blocks_by_page for block in blocks]
        pages = [ParsedPage(None, "\n\n".join(all_blocks))] if all_blocks else []
        page_count = None
    if not pages:
        raise ParseError("The document contains no text.")
    return ParsedDocument(pages, page_count)


def _check_docx_size(data: BinaryIO) -> None:
    # Guards against zip bombs: a tiny archive that expands to gigabytes in the worker.
    with zipfile.ZipFile(data) as archive:
        total = sum(info.file_size for info in archive.infolist())
    data.seek(0)
    if total > MAX_DOCX_UNCOMPRESSED_BYTES:
        raise ParseError("The DOCX is too large when unpacked.")


def _read_docx_body(document) -> tuple[list[list[str]], bool]:
    pages: list[list[str]] = [[]]
    saw_rendered_break = False

    def start_new_page() -> None:
        # No-op on an empty page: Word writes both a hard break and a rendered break at the
        # same spot, and we must not count that boundary twice.
        if pages[-1]:
            pages.append([])

    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            fragment: list[str] = []

            def flush(fragment: list[str] = fragment) -> None:
                text = "".join(fragment).strip()
                if text:
                    pages[-1].append(text)
                fragment.clear()

            for run in element.iter(qn("w:r")):
                for child in run:
                    if child.tag == qn("w:t"):
                        fragment.append(child.text or "")
                    elif child.tag in (qn("w:tab"), qn("w:noBreakHyphen")):
                        fragment.append("\t" if child.tag == qn("w:tab") else "-")
                    elif child.tag == qn("w:cr"):
                        fragment.append("\n")
                    elif child.tag == qn("w:br"):
                        if child.get(qn("w:type")) == "page":
                            flush()
                            start_new_page()
                        else:
                            fragment.append("\n")
                    elif child.tag == qn("w:lastRenderedPageBreak"):
                        saw_rendered_break = True
                        flush()
                        start_new_page()
            flush()
        elif element.tag == qn("w:tbl"):
            text = _table_text(Table(element, document))
            if text:
                pages[-1].append(text)
    return pages, saw_rendered_break


def _table_text(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells: list[str] = []
        previous = None
        for cell in row.cells:
            if cell._tc is previous:  # merged cells repeat the same underlying element
                continue
            previous = cell._tc
            if cell.text.strip():
                cells.append(cell.text.strip())
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def decode_text(raw: bytes) -> str:
    # cp1250 is the usual encoding of older Serbian/Central European text files from Windows.
    for encoding in ("utf-8-sig", "cp1250"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("The text file uses an unsupported encoding. Save it as UTF-8 and try again.")


def parse_text_file(data: BinaryIO) -> ParsedDocument:
    return parse_plain_text(decode_text(data.read()))


def parse_plain_text(text: str) -> ParsedDocument:
    cleaned = _normalize_newlines(text)
    if not cleaned:
        raise ParseError("The text is empty.")
    return ParsedDocument([ParsedPage(None, cleaned)], None)
