"""Builders for in-memory test files (PDF, DOCX, ...), so the repo holds no binary fixtures."""

import io
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from fpdf import FPDF
from PIL import Image
from pypdf import PdfReader, PdfWriter

SERBIAN_FONT = Path("/usr/share/fonts/liberation-sans-fonts/LiberationSans-Regular.ttf")

LOREM = (
    "The quick brown fox jumps over the lazy dog while the students read their notes carefully. "
    "Every page of this document contains enough ordinary text to look like a real lecture."
)


def make_pdf(pages: list[str], serbian_font: bool = False) -> io.BytesIO:
    pdf = FPDF()
    if serbian_font:
        pdf.add_font("Lib", "", str(SERBIAN_FONT))
    for text in pages:
        pdf.add_page()
        pdf.set_font("Lib" if serbian_font else "Helvetica", size=12)
        if text:
            pdf.multi_cell(0, 8, text)
    return io.BytesIO(bytes(pdf.output()))


def make_scanned_pdf(page_count: int = 3) -> io.BytesIO:
    image = Image.new("RGB", (600, 800), "white")
    buffer = io.BytesIO()
    image.save(
        buffer, "PDF", save_all=True, append_images=[image.copy() for _ in range(page_count - 1)]
    )
    buffer.seek(0)
    return buffer


def encrypt_pdf(source: io.BytesIO, user_password: str, owner_password: str) -> io.BytesIO:
    writer = PdfWriter()
    writer.append(PdfReader(source))
    writer.encrypt(user_password=user_password, owner_password=owner_password)
    buffer = io.BytesIO()
    writer.write(buffer)
    buffer.seek(0)
    return buffer


def make_docx(paragraphs: list[str | None], table: list[list[str]] | None = None) -> Document:
    """A None entry becomes a hard page break."""
    document = Document()
    for text in paragraphs:
        if text is None:
            document.add_page_break()
        else:
            document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    return document


def add_rendered_break(document: Document, before_paragraph_index: int) -> None:
    """Mimics what Word saves: a marker where it laid out a new page."""
    paragraph = document.paragraphs[before_paragraph_index]
    run = paragraph.runs[0]._r
    run.insert(0, OxmlElement("w:lastRenderedPageBreak"))


def docx_bytes(document: Document) -> io.BytesIO:
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def make_zip_bomb(megabytes: int = 101) -> io.BytesIO:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", b"\0" * (megabytes * 1024 * 1024))
    buffer.seek(0)
    return buffer
